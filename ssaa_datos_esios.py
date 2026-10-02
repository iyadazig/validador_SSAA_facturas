# -*- coding: utf-8 -*-
"""
LECTURA DE LOS EXCEL HISTORICOS DE ESIOS
========================================
Lee (sin modificarlos) los Excel que genera el proyecto Descarga_datos_ESIOS:

  AAAA_Historico_PVPC_horario.xlsx      Total SAH y PERD del PVPC, horario
  AAAA_Historico_componentes_ESIOS.xlsx SSAA cuartohorario por componente y
                                        pestanas "Perdidas mmm-aa" / "Perdidas QH mmm-aa"

Todas las series se devuelven como {clave: valor} con la convencion de ESIOS:
  horaria       (fecha, hora)          hora 1..23/24/25 segun el dia
  cuartohoraria (fecha, hora, cuarto)  cuarto 1..4
Asi los dias de cambio de hora no necesitan zona horaria.
"""

import datetime as dt
from collections import defaultdict

from openpyxl import load_workbook

import config

MESES = ["ene", "feb", "mar", "abr", "may", "jun",
         "jul", "ago", "sep", "oct", "nov", "dic"]

COL_TOTAL_SSAA = "TOTAL SSAA EUR/MWh (suma)"
# Columnas del Excel de componentes que son SSAA (seleccionables como "componentes").
PREFIJOS_SSAA = ("RT3", "RT6", "BS3", "DSV", "EXD", "IN7", "BALX", "CT3", "CFP",
                 "SRAD", "CT2")

_cache = {}


class SinDatos(Exception):
    """El periodo pedido no esta en los Excel historicos."""


def nombre_mes(fecha):
    return "%s-%02d" % (MESES[fecha.month - 1], fecha.year % 100)


def meses_entre(ini, fin):
    m = dt.date(ini.year, ini.month, 1)
    while m <= fin:
        yield m
        m = dt.date(m.year + (m.month == 12), m.month % 12 + 1, 1)


def dias_entre(ini, fin):
    d = ini
    while d <= fin:
        yield d
        d += dt.timedelta(days=1)


def _fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _filas(ruta, hoja):
    """Filas de una hoja como tuplas, cacheadas por fecha de modificacion."""
    if not ruta.exists():
        raise SinDatos("No existe %s" % ruta)
    clave = (str(ruta), hoja, ruta.stat().st_mtime)
    if clave not in _cache:
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            if hoja not in wb.sheetnames:
                _cache[clave] = None
            else:
                _cache[clave] = [tuple(r) for r in wb[hoja].iter_rows(values_only=True)]
        finally:
            wb.close()
    return _cache[clave]


def _ruta_pvpc(anio):
    return config.CARPETA_ESIOS / ("%d_Historico_PVPC_horario.xlsx" % anio)


def _ruta_componentes(anio):
    return config.CARPETA_ESIOS / ("%d_Historico_componentes_ESIOS.xlsx" % anio)


def _buscar_col(cabecera, texto, desde=0, hasta=None):
    hasta = len(cabecera) if hasta is None else hasta
    for i in range(desde, hasta):
        c = cabecera[i]
        if c is not None and texto.lower() in str(c).lower():
            return i
    raise SinDatos("No encuentro la columna '%s'" % texto)


# ------------------------------------------------------------------ PVPC horario
def pvpc_horario(ini, fin):
    """{(fecha, hora): {'sah': EUR/MWh bc, 'perd': %}} del PVPC_DETALLE."""
    out = {}
    for m in meses_entre(ini, fin):
        filas = _filas(_ruta_pvpc(m.year), nombre_mes(m))
        if not filas:
            continue
        cab = filas[0]
        c_sah = _buscar_col(cab, "Total SAH")
        c_perd = _buscar_col(cab, "PVPC PERD")
        c_hora = _buscar_col(cab, "Hora", 1)
        for r in filas[1:]:
            f = _fecha(r[0])
            if f is None or not (ini <= f <= fin) or r[c_sah] is None:
                continue
            out[(f, int(r[c_hora]))] = {"sah": float(r[c_sah]),
                                        "perd": float(r[c_perd] or 0)}
    return out


# ------------------------------------------------- componentes cuartohorarios
def columnas_componentes(anio=None):
    """Nombres de las columnas de SSAA seleccionables del Excel de componentes."""
    anio = anio or dt.date.today().year
    ruta = _ruta_componentes(anio)
    wb = load_workbook(ruta, read_only=True)
    try:
        cab = next(wb[wb.sheetnames[0]].iter_rows(min_row=3, max_row=3, values_only=True))
    finally:
        wb.close()
    return [c for c in cab if c and (str(c).startswith(PREFIJOS_SSAA) or c == COL_TOTAL_SSAA)]


def componentes_qh(ini, fin, columnas=(COL_TOTAL_SSAA,)):
    """{(fecha, hora, cuarto): suma de las columnas pedidas} y {mes: liquidacion}."""
    out, liquidaciones = {}, {}
    for m in meses_entre(ini, fin):
        filas = _filas(_ruta_componentes(m.year), nombre_mes(m))
        if not filas:
            continue
        cab = filas[2]
        idx = [_buscar_col(cab, c) for c in columnas]
        for r in filas[3:]:
            f = _fecha(r[0])
            if f is None or not (ini <= f <= fin):
                continue
            vals = [r[i] for i in idx]
            if all(v is None for v in vals):
                continue
            out[(f, int(r[1]), int(r[2]))] = sum(float(v or 0) for v in vals)
            liquidaciones[nombre_mes(m)] = r[-1]
    return out, liquidaciones


# ------------------------------------------------------------------- perdidas
def perdidas_horarias(ini, fin, tarifa, zona="Península"):
    """{(fecha, hora): PERD %} del liquicomun (pestanas 'Perdidas mmm-aa')."""
    out = {}
    for m in meses_entre(ini, fin):
        filas = _filas(_ruta_componentes(m.year), "Pérdidas " + nombre_mes(m))
        if not filas:
            continue
        zonas, cab = filas[1], filas[2]
        # la fila de zonas solo tiene texto en la primera columna de cada bloque
        desde = next((i for i, z in enumerate(zonas) if z and zona.lower() in str(z).lower()), None)
        if desde is None:
            raise SinDatos("Zona '%s' no encontrada en Pérdidas %s" % (zona, nombre_mes(m)))
        hasta = next((i for i in range(desde + 1, len(zonas)) if zonas[i]), len(zonas))
        col = _buscar_col(cab, "PERD %s %%" % tarifa, desde, hasta)
        for r in filas[3:]:
            f = _fecha(r[0])
            if f is None or not (ini <= f <= fin) or r[col] is None:
                continue
            out[(f, int(r[1]))] = float(r[col])
    return out


def perdidas_qh(ini, fin, tarifa):
    """{(fecha, hora, cuarto): PERD %} peninsular (pestanas 'Perdidas QH mmm-aa')."""
    out = {}
    for m in meses_entre(ini, fin):
        filas = _filas(_ruta_componentes(m.year), "Pérdidas QH " + nombre_mes(m))
        if not filas:
            continue
        col = _buscar_col(filas[1], "PERD QH %s %%" % tarifa)
        for r in filas[2:]:
            f = _fecha(r[0])
            if f is None or not (ini <= f <= fin) or r[col] is None:
                continue
            out[(f, int(r[1]), int(r[2]))] = float(r[col])
    return out


# ------------------------------------------------------------------ utilidades
def a_horario(serie_qh):
    """Media de los cuartos de cada hora."""
    acum = defaultdict(list)
    for (f, h, _q), v in serie_qh.items():
        acum[(f, h)].append(v)
    return {k: sum(v) / len(v) for k, v in acum.items()}


def dias_sin_datos(serie, ini, fin):
    """Dias del periodo sin ninguna hora en la serie."""
    con = {k[0] for k in serie}
    return [d for d in dias_entre(ini, fin) if d not in con]
