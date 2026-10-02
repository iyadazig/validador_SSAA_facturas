# -*- coding: utf-8 -*-
"""Informe Excel de una revision y fichas de contrato guardadas por CUPS."""

import datetime as dt
import io
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import config
from ssaa_motor import (Contrato, INDICES, AGREGACIONES, MECANISMOS, PERDIDAS,
                        PERD_AGREGACIONES)

NEGRITA = Font(bold=True)
CABECERA = PatternFill("solid", fgColor="DDEBF7")
COLORES = {"CORRECTO": "C6EFCE", "FACTURADO DE MÁS": "FFC7CE", "FACTURADO DE MENOS": "FFEB9C"}


# ------------------------------------------------------------------ contratos
def cargar_contratos():
    if not config.FICHERO_CONTRATOS.exists():
        return {}
    with open(config.FICHERO_CONTRATOS, encoding="utf-8") as f:
        return {k: Contrato.desde_dict(v) for k, v in json.load(f).items()}


def guardar_contrato(contrato):
    if not contrato.cups:
        raise ValueError("El contrato necesita CUPS para guardarse")
    todos = {k: v.a_dict() for k, v in cargar_contratos().items()}
    todos[contrato.cups] = contrato.a_dict()
    tmp = str(config.FICHERO_CONTRATOS) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(todos, f, ensure_ascii=False, indent=2, sort_keys=True)
    os.replace(tmp, config.FICHERO_CONTRATOS)


# -------------------------------------------------------------------- informe
def _tabla(ws, fila, filas):
    for i, (k, v) in enumerate(filas):
        ws.cell(fila + i, 1, k).font = NEGRITA
        ws.cell(fila + i, 2, v)
    return fila + len(filas) + 1


def _titulo(ws, fila, texto):
    c = ws.cell(fila, 1, texto)
    c.font = Font(bold=True, size=12)
    return fila + 1


def descripcion_contrato(c):
    filas = [("Tarifa", c.tarifa), ("Índice", INDICES[c.indice])]
    if c.indice == "componentes":
        filas.append(("Componentes", "; ".join(c.componentes)))
    filas += [("Agregación", AGREGACIONES[c.agregacion]),
              ("Mecanismo", MECANISMOS[c.mecanismo])]
    if c.mecanismo == "banda":
        filas += [("Ref. superior €/MWh", c.ref_superior),
                  ("Ref. inferior €/MWh", c.ref_inferior)]
    if c.mecanismo in ("techo", "suelo_techo"):
        filas.append(("Techo €/MWh", c.techo))
    if c.mecanismo == "suelo_techo":
        filas.append(("Suelo €/MWh", c.suelo))
    if c.mecanismo == "fijo":
        filas.append(("Precio fijo €/MWh", c.precio_fijo))
    if c.mecanismo in ("indexado", "techo", "suelo_techo") and c.prima:
        filas.append(("Prima €/MWh", c.prima))
    filas.append(("Pérdidas", PERDIDAS[c.perdidas]))
    if c.perdidas == "fijo":
        filas.append(("Coef. pérdidas %", c.perd_fijo))
    elif c.perdidas != "ninguna" and c.agregacion != "horaria":
        filas.append(("Agregación pérdidas", PERD_AGREGACIONES[c.perd_agregacion]))
    filas.append(("Factor final", c.factor))
    return filas


def generar_excel(r, factura, diagnostico=None):
    """Excel de la revision en bytes. `factura` es un dict de campos a mostrar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    fila = _titulo(ws, 1, "Revisión de SSAA — %s" % dt.date.today().strftime("%d/%m/%Y"))
    fila = _titulo(ws, fila + 1, "Factura")
    fila = _tabla(ws, fila, list(factura.items()))
    fila = _titulo(ws, fila, "Cláusula del contrato")
    fila = _tabla(ws, fila, descripcion_contrato(r.contrato))
    fila = _titulo(ws, fila, "Resultado")
    filas = [("Periodo", "%s a %s" % (r.inicio.strftime("%d/%m/%Y"), r.fin.strftime("%d/%m/%Y"))),
             ("Energía MWh", round(r.energia_mwh, 6)),
             ("Índice medio aritmético €/MWh", round(r.indice_medio, 6))]
    if r.indice_ponderado is not None:
        filas.append(("Índice medio ponderado €/MWh", round(r.indice_ponderado, 6)))
    filas += [("Precio aplicado €/MWh (antes de pérdidas)", round(r.precio_aplicado, 6)),
              ("Pérdidas aplicadas %", round(r.perd_aplicada, 4)),
              ("Importe recalculado €", round(r.importe, 2)),
              ("Importe facturado €", r.facturado),
              ("Diferencia € (facturado − recalculado)",
               round(r.diferencia, 2) if r.diferencia is not None else None),
              ("Diferencia %", round(r.diferencia_pct, 3) if r.diferencia_pct is not None else None),
              ("Veredicto", r.veredicto)]
    fila_ver = fila + len(filas) - 1
    fila = _tabla(ws, fila, filas)
    if r.veredicto in COLORES:
        ws.cell(fila_ver, 2).fill = PatternFill("solid", fgColor=COLORES[r.veredicto])
        ws.cell(fila_ver, 2).font = NEGRITA
    fila = _titulo(ws, fila, "Datos ESIOS usados")
    fila = _tabla(ws, fila, list(r.liquidaciones.items()) or [("—", "")])
    if r.avisos:
        fila = _titulo(ws, fila, "Avisos")
        for a in r.avisos:
            ws.cell(fila, 1, a)
            fila += 1
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 60

    # detalle hora a hora
    wd = wb.create_sheet("Detalle")
    cab = ["Fecha", "Hora"] + (["Cuarto"] if r.resolucion == "qh" else []) + \
          ["Consumo kWh", "Índice €/MWh", "Pérdidas %"]
    if r.contrato.agregacion == "horaria":
        cab += ["Precio aplicado €/MWh", "Importe €"]
    wd.append(cab)
    for k, e, v, p, precio, imp in r.detalle:
        fila_d = [k[0], k[1]] + ([k[2]] if r.resolucion == "qh" else []) + [e, v, p]
        if r.contrato.agregacion == "horaria":
            fila_d += [precio, imp]
        wd.append(fila_d)
    for i, c in enumerate(cab, 1):
        wd.cell(1, i).font = NEGRITA
        wd.cell(1, i).fill = CABECERA
        wd.cell(1, i).alignment = Alignment(wrap_text=True)
        wd.column_dimensions[get_column_letter(i)].width = 14
    for celda in wd["A"][1:]:
        celda.number_format = "dd/mm/yyyy"
    wd.freeze_panes = "A2"

    if diagnostico:
        wg = wb.create_sheet("Diagnóstico")
        cab = list(diagnostico[0].keys())
        wg.append(cab)
        for d in diagnostico:
            wg.append([d[c] for c in cab])
        for i in range(1, len(cab) + 1):
            wg.cell(1, i).font = NEGRITA
            wg.cell(1, i).fill = CABECERA
            wg.column_dimensions[get_column_letter(i)].width = 26

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def guardar_informe(datos, r, cups, numero):
    config.CARPETA_REVISIONES.mkdir(exist_ok=True)
    limpio = lambda s: re.sub(r"[^A-Za-z0-9_-]", "", s or "") or "sin"
    nombre = "%s_%s_%s.xlsx" % (r.inicio.strftime("%Y-%m"), limpio(cups), limpio(numero))
    ruta = config.CARPETA_REVISIONES / nombre
    ruta.write_bytes(datos)
    return ruta
