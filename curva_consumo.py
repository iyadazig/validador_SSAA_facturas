# -*- coding: utf-8 -*-
"""
CURVA DE CONSUMO
================
Lee una curva horaria o cuartohoraria de un CSV/XLSX/XLS (exportacion de
Gemweb, Datadis o distribuidora) y la deja en la convencion de ESIOS:
  horaria       {(fecha, hora 1..25): kWh}
  cuartohoraria {(fecha, hora, cuarto 1..4): kWh}

Las horas se numeran por orden dentro de cada dia, asi los dias de cambio de
hora (23 o 25 registros) cuadran con ESIOS sin manejar zonas horarias. Si el
fichero etiqueta cada registro con la hora FINAL (01:00 ... 24:00/00:00 del dia
siguiente) se detecta y se corrige.
Si un dia de cambio de hora trae 24 h de reloj (96 cuartos), como Gemweb, se
reparte con curva_reloj_a_esios(). La API de Gemweb esta en gemweb.py.
"""

import csv
import datetime as dt
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field


PALABRAS_FECHA = ("fecha", "date", "timestamp")
PALABRAS_HORA = ("hora", "hour", "periodo", "time")
PALABRAS_KWH = ("kwh", "consumo", "activa", "energia", "energía", "ae", "valor", "value")
EXCLUIR_KWH = ("reactiv", "export", "excedent", "gener", "r1", "r2", "r3", "r4",
               "metodo", "método", "cups", "calidad", "potencia")


@dataclass
class Curva:
    resolucion: str                       # "h" o "qh"
    valores: dict                         # clave ESIOS -> kWh
    columnas: dict = field(default_factory=dict)
    avisos: list = field(default_factory=list)
    origen: str = ""

    @property
    def total_kwh(self):
        return sum(self.valores.values())

    @property
    def inicio(self):
        return min(k[0] for k in self.valores) if self.valores else None

    @property
    def fin(self):
        return max(k[0] for k in self.valores) if self.valores else None


class ErrorCurva(Exception):
    pass


# --------------------------------------------------------------- lectura bruta
def _filas_csv(datos):
    for cod in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = datos.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:5000]
    sep = max((";", ",", "\t", "|"), key=muestra.count)
    return [r for r in csv.reader(io.StringIO(texto), delimiter=sep)]


def _filas_xlsx(datos):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _filas_xls(datos):
    import xlrd
    libro = xlrd.open_workbook(file_contents=datos)
    hoja = libro.sheet_by_index(0)
    filas = []
    for i in range(hoja.nrows):
        fila = []
        for c in hoja.row(i):
            if c.ctype == xlrd.XL_CELL_DATE:
                fila.append(xlrd.xldate.xldate_as_datetime(c.value, libro.datemode))
            else:
                fila.append(c.value)
        filas.append(fila)
    return filas


def filas_fichero(datos, nombre):
    n = nombre.lower()
    if n.endswith(".xlsx") or n.endswith(".xlsm"):
        return _filas_xlsx(datos)
    if n.endswith(".xls"):
        return _filas_xls(datos)
    return _filas_csv(datos)


# ------------------------------------------------------------------- parseo
def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") \
            else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


_FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y", "%d/%m/%y")


def _fecha_hora(v):
    """Devuelve (datetime o date, tiene_hora)."""
    if isinstance(v, dt.datetime):
        return v, True
    if isinstance(v, dt.date):
        return v, False
    s = str(v or "").strip()
    if not s:
        return None, False
    s = s.replace("T", " ").split("+")[0].rstrip("Z")
    partes = s.split()
    for f in _FORMATOS_FECHA:
        try:
            d = dt.datetime.strptime(partes[0], f)
        except ValueError:
            continue
        if len(partes) > 1:
            hm = _hora_minutos(partes[1])
            if hm is not None:
                return d + dt.timedelta(minutes=hm), True
        return d.date(), False
    return None, False


def _hora_minutos(v):
    """'01:00', '1', 1, '24:00', time -> minutos desde las 00:00 (o None)."""
    if isinstance(v, dt.time):
        return v.hour * 60 + v.minute
    if isinstance(v, dt.datetime):
        return v.hour * 60 + v.minute
    if isinstance(v, (int, float)):
        return None
    m = re.match(r"^(\d{1,2})[:.h](\d{2})", str(v).strip())
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    return None


def _detectar_cabecera(filas):
    for i, fila in enumerate(filas[:20]):
        textos = [str(c).strip().lower() for c in fila if isinstance(c, str) and c.strip()]
        if len(textos) >= 2 and any(any(p in t for p in PALABRAS_FECHA)
                                    or t in ("dia", "día", "day") for t in textos):
            return i
    raise ErrorCurva("No encuentro la fila de cabecera (columna de fecha)")


def detectar_columnas(cabecera):
    cab = [str(c or "").strip().lower() for c in cabecera]
    fecha = next((i for i, c in enumerate(cab) if any(p in c for p in PALABRAS_FECHA)),
                 next((i for i, c in enumerate(cab) if c in ("dia", "día", "day")), None))
    hora = next((i for i, c in enumerate(cab) if i != fecha
                 and any(p in c for p in PALABRAS_HORA)), None)
    candidatas = [i for i, c in enumerate(cab) if i not in (fecha, hora) and c
                  and any(p in c for p in PALABRAS_KWH)
                  and not any(x in c for x in EXCLUIR_KWH)]
    # preferir la que diga kWh explicitamente
    candidatas.sort(key=lambda i: ("kwh" not in cab[i], "activa" not in cab[i]
                                   and "consumo" not in cab[i]))
    return {"fecha": fecha, "hora": hora, "kwh": candidatas[0] if candidatas else None}


def _factor_unidad(nombre):
    n = nombre.lower()
    if "mwh" in n:
        return 1000.0
    if "kwh" in n:
        return 1.0
    if "wh" in n:
        return 0.001
    return 1.0


def leer_curva(datos, nombre, columnas=None):
    """Lee una curva de un fichero (bytes). `columnas` permite forzar
    {'fecha': i, 'hora': i|None, 'kwh': i}."""
    filas = filas_fichero(datos, nombre)
    ic = _detectar_cabecera(filas)
    cabecera = filas[ic]
    cols = dict(detectar_columnas(cabecera))
    if columnas:
        cols.update({k: v for k, v in columnas.items() if k in cols})
    if cols["fecha"] is None or cols["kwh"] is None:
        raise ErrorCurva("No identifico las columnas de fecha y consumo: %s" % cabecera)
    factor = _factor_unidad(str(cabecera[cols["kwh"]]))

    registros = []      # (instante o (fecha, hora_int), orden, kWh)
    con_hora_entera = False
    for n, fila in enumerate(filas[ic + 1:]):
        if len(fila) <= max(c for c in cols.values() if c is not None):
            continue
        kwh = _num(fila[cols["kwh"]])
        f, tiene_hora = _fecha_hora(fila[cols["fecha"]])
        if f is None or kwh is None:
            continue
        if not tiene_hora and cols["hora"] is not None:
            v = fila[cols["hora"]]
            minutos = _hora_minutos(v)
            if minutos is not None:
                f = dt.datetime.combine(f, dt.time()) + dt.timedelta(minutes=minutos)
                tiene_hora = True
            else:
                h = _num(v)
                if h is None:
                    continue
                registros.append(((f, int(h)), n, kwh * factor))
                con_hora_entera = True
                continue
        if not tiene_hora:
            raise ErrorCurva("La columna de fecha no trae hora y no hay columna de hora")
        registros.append((f, n, kwh * factor))

    if not registros:
        raise ErrorCurva("El fichero no tiene registros de consumo legibles")

    avisos = []
    por_dia = defaultdict(list)
    if con_hora_entera:
        # columna de hora numerica: 1..24/25 (o 0..23)
        base = 1 if min(r[0][1] for r in registros) == 0 else 0
        for (f, h), n, kwh in registros:
            por_dia[f].append((h + base, n, kwh))
    else:
        instantes = sorted(r[0] for r in registros)
        paso = min((b - a for a, b in zip(instantes, instantes[1:]) if b > a),
                   default=dt.timedelta(hours=1))
        primero = instantes[0]
        # etiquetado con la hora final si el primer registro no cae a las 00:00
        fin_intervalo = (primero.hour, primero.minute) != (0, 0)
        if fin_intervalo:
            avisos.append("Registros etiquetados con la hora final del intervalo; "
                          "se desplazan %d min." % (paso.total_seconds() / 60))
        for t, n, kwh in registros:
            if fin_intervalo:
                t = t - paso
            por_dia[t.date()].append((t, n, kwh))

    # resolucion por numero mediano de registros al dia
    tamanos = sorted(len(v) for v in por_dia.values())
    resolucion = "qh" if tamanos[len(tamanos) // 2] > 30 else "h"

    valores = {}
    nominal = 96 if resolucion == "qh" else 24
    paso_min = 15 if resolucion == "qh" else 60
    for f, regs in por_dia.items():
        regs.sort(key=lambda r: (r[0], r[1]))
        if f in dias_cambio_hora(f.year) and len(regs) == nominal:
            # dia de cambio de hora normalizado a 24 h de reloj: se reparte por hora real
            reloj = [(f, (r[0].hour * 60 + r[0].minute) if isinstance(r[0], dt.datetime)
                      else (pos - 1) * paso_min, r[2]) for pos, r in enumerate(regs, 1)]
            parcial, av = curva_reloj_a_esios(reloj, resolucion)
            for k, v in parcial.items():
                valores[k] = valores.get(k, 0.0) + v
            avisos.extend(av)
            continue
        for pos, (_t, _n, kwh) in enumerate(regs, 1):
            if resolucion == "qh":
                clave = (f, (pos - 1) // 4 + 1, (pos - 1) % 4 + 1)
            else:
                clave = (f, pos)
            valores[clave] = valores.get(clave, 0.0) + kwh
        esperado = (92, 96, 100) if resolucion == "qh" else (23, 24, 25)
        if len(regs) not in esperado:
            avisos.append("%s: %d registros (incompleto)." % (f, len(regs)))

    return Curva(resolucion=resolucion, valores=valores,
                 columnas={k: (str(cabecera[v]) if v is not None else None)
                           for k, v in cols.items()},
                 avisos=avisos, origen=nombre)


# ------------------------------------------------------------ cambio de hora
def _ultimo_domingo(anio, mes):
    d = dt.date(anio + (mes == 12), mes % 12 + 1, 1) - dt.timedelta(days=1)
    return d - dt.timedelta(days=(d.weekday() + 1) % 7)


def dias_cambio_hora(anio):
    """{dia de marzo (23 h), dia de octubre (25 h)} en Espana peninsular."""
    return {_ultimo_domingo(anio, 3), _ultimo_domingo(anio, 10)}


def curva_reloj_a_esios(registros, resolucion):
    """[(fecha, minuto de inicio en hora de reloj, kWh)] -> ({clave ESIOS: kWh}, avisos).

    Para curvas que traen siempre 24 h de reloj (Gemweb da 96 cuartos/dia):
      marzo   ESIOS hora 1-2 = 00-02, 3-23 = 03-24. Lo que venga en las 02:xx
              (hora que no existe) se suma a la hora 3.
      octubre ESIOS hora 1-2 = 00-02, 3 y 4 = las dos 02-03, 5-25 = 03-24.
              Lo de las 02:xx se reparte a medias entre las horas 3 y 4.
    """
    out, avisos = {}, []
    inexistente = defaultdict(float)
    repetida = defaultdict(float)
    for f, minuto, kwh in registros:
        h0, q = minuto // 60, (minuto % 60) // 15 + 1
        marzo, octubre = sorted(dias_cambio_hora(f.year))
        if f == marzo and h0 >= 2:
            horas = [(3 if h0 == 2 else h0, 1.0)]
            if h0 == 2:
                inexistente[f] += kwh
        elif f == octubre and h0 >= 2:
            horas = [(3, 0.5), (4, 0.5)] if h0 == 2 else [(h0 + 2, 1.0)]
            if h0 == 2:
                repetida[f] += kwh
        else:
            horas = [(h0 + 1, 1.0)]
        for h, frac in horas:
            k = (f, h, q) if resolucion == "qh" else (f, h)
            out[k] = out.get(k, 0.0) + kwh * frac
    for f, e in inexistente.items():
        if e:
            avisos.append("%s (cambio a horario de verano): %.3f kWh en las 02:00-03:00, "
                          "hora que no existe; sumados a la hora siguiente." % (f, e))
    for f, e in repetida.items():
        avisos.append("%s (cambio a horario de invierno): la curva trae una sola hora "
                      "02:00-03:00 (%.3f kWh); se reparte a medias entre las dos horas "
                      "de ESIOS." % (f, e))
    return out, avisos
